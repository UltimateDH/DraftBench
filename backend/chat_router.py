import base64
import os
from typing import List, Literal

import anthropic
from fastapi import APIRouter, Depends, HTTPException
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

import models, database, auth_utils
from materials_router import UPLOAD_DIR

router = APIRouter()

# "anthropic" (Claude API) or "ollama" (local model, e.g. qwen2.5:7b-instruct)
PROVIDER = os.getenv("AI_PROVIDER", "anthropic").lower()

ANTHROPIC_MODEL = os.getenv("ANTHROPIC_MODEL", "claude-sonnet-5-5")
MAX_TEXT_CHARS = 300_000  # Claude path: total extracted text per question

OLLAMA_MODEL = os.getenv("OLLAMA_MODEL", "qwen2.5:7b-instruct")
OLLAMA_NUM_CTX = int(os.getenv("OLLAMA_NUM_CTX", "8192"))
OLLAMA_MAX_CONTEXT_CHARS = int(os.getenv("OLLAMA_MAX_CONTEXT_CHARS", "10000"))  # Devanagari uses many tokens
OLLAMA_NUM_PREDICT = int(os.getenv("OLLAMA_NUM_PREDICT", "600"))

IMAGE_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp"}

SYSTEM = (
    "You are the Draft Bench assistant, helping a teacher in Nepal with the materials they uploaded. "
    "Answer using only the attached materials. If the answer isn't in them, say so instead of guessing. "
    "Mention which material (and the page or slide, when you can tell) your answer comes from. "
    "Reply in the language the teacher writes in. When writing Nepali, use clear, natural Devanagari Nepali, "
    "do not use Hindi vocabulary or Bengali characters. "
    "Write in plain text only: no markdown, no asterisks, no pound signs. "
    "For lists, use numbered lines or lines starting with a hyphen. "
    "Keep answers clear and concise unless the teacher asks for more detail."
)


class Turn(BaseModel):
    role: Literal["user", "assistant"]
    content: str


class ChatRequest(BaseModel):
    material_ids: List[int] = Field(min_length=1, max_length=10)
    message: str = Field(min_length=1, max_length=4000)
    history: List[Turn] = []


class ChatResponse(BaseModel):
    reply: str


def extract_text(path: str, ext: str) -> str:
    if ext in (".txt", ".md", ".csv"):
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            return f.read()
    if ext == ".pdf":  # only used by the Ollama path; Claude reads PDFs natively
        from pypdf import PdfReader
        return "\n".join(
            f"[Page {i}]\n{page.extract_text() or ''}" for i, page in enumerate(PdfReader(path).pages, 1)
        )
    if ext == ".docx":
        import docx
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if ext == ".pptx":
        from pptx import Presentation
        parts = []
        for i, slide in enumerate(Presentation(path).slides, 1):
            parts.append(f"[Slide {i}]")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text_frame.text)
        return "\n".join(parts)
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"[Sheet {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                parts.append(" | ".join("" if c is None else str(c) for c in row))
        return "\n".join(parts)
    raise ValueError(ext)


def material_path(m: models.Material):
    path = os.path.join(UPLOAD_DIR, m.stored_name)
    if not os.path.exists(path):
        raise HTTPException(status_code=410, detail=f'"{m.filename}" is missing on the server. Re-upload it.')
    return path, os.path.splitext(m.stored_name)[1].lower()


def read_text(m: models.Material) -> str:
    path, ext = material_path(m)
    try:
        return extract_text(path, ext)
    except ValueError:
        raise HTTPException(
            status_code=400,
            detail=f'Chat can\'t read "{m.filename}" yet. Save it as PDF or DOCX and upload that instead.',
        )


# ---------- Claude (reads PDFs and images natively) ----------
def build_blocks(materials: List[models.Material]) -> list:
    blocks, total = [], 0
    for m in materials:
        path, ext = material_path(m)
        blocks.append({"type": "text", "text": f'Material "{m.filename}":'})
        if ext == ".pdf" or ext in IMAGE_TYPES:
            with open(path, "rb") as f:
                data = base64.standard_b64encode(f.read()).decode()
            if ext == ".pdf":
                blocks.append({"type": "document", "source": {"type": "base64", "media_type": "application/pdf", "data": data}})
            else:
                blocks.append({"type": "image", "source": {"type": "base64", "media_type": IMAGE_TYPES[ext], "data": data}})
        else:
            text = read_text(m)
            total += len(text)
            if total > MAX_TEXT_CHARS:
                raise HTTPException(status_code=413, detail="These materials are too long to chat with at once. Select fewer files.")
            blocks.append({"type": "text", "text": text})
    return blocks


def ask_anthropic(materials, msgs) -> str:
    if not os.getenv("ANTHROPIC_API_KEY"):
        raise HTTPException(status_code=500, detail="AI isn't configured on the server (missing ANTHROPIC_API_KEY).")
    msgs[0]["content"] = build_blocks(materials) + [{"type": "text", "text": msgs[0]["content"]}]
    try:
        resp = anthropic.Anthropic().messages.create(model=ANTHROPIC_MODEL, max_tokens=1500, system=SYSTEM, messages=msgs)
    except anthropic.APIStatusError as e:
        raise HTTPException(status_code=502, detail=f"AI service error: {e.message}")
    except anthropic.APIConnectionError:
        raise HTTPException(status_code=502, detail="Couldn't reach the AI service. Try again.")
    return "".join(b.text for b in resp.content if b.type == "text")


# ---------- Ollama (local model, text only) ----------
def ask_ollama(materials, msgs) -> str:
    import ollama

    parts, used, truncated = [], 0, False
    for m in materials:
        _, ext = material_path(m)
        if ext in IMAGE_TYPES:
            raise HTTPException(status_code=400, detail=f'The local model can\'t read images ("{m.filename}"). Upload a PDF or DOCX instead.')
        text = read_text(m).strip()
        room = OLLAMA_MAX_CONTEXT_CHARS - used
        if room <= 0:
            truncated = True
            break
        if len(text) > room:
            text, truncated = text[:room], True
        used += len(text)
        parts.append(f'Material "{m.filename}":\n{text}')
    context = "\n\n".join(parts)
    if truncated:
        context += "\n\n[The materials were cut short because they are long. Tell the teacher if the answer may be in the missing part.]"
    if not context.strip():
        raise HTTPException(status_code=400, detail="Couldn't find any readable text in these files. Scanned PDFs need the Claude provider.")

    msgs[0]["content"] = f"{context}\n\n---\nTeacher's question: {msgs[0]['content']}"
    try:
        resp = ollama.Client(host=os.getenv("OLLAMA_HOST")).chat(
            model=OLLAMA_MODEL,
            messages=[{"role": "system", "content": SYSTEM}] + msgs,
            options={
                "temperature": 0.2,
                "top_p": 0.85,
                "repeat_penalty": 1.1,
                "num_ctx": OLLAMA_NUM_CTX,
                "num_predict": OLLAMA_NUM_PREDICT,
            },
            stream=False,
        )
    except ollama.ResponseError as e:
        raise HTTPException(status_code=502, detail=f"Local model error: {e.error}")
    except (ConnectionError, OSError):
        raise HTTPException(status_code=502, detail="Couldn't reach Ollama. Is it running on the server?")
    return resp["message"]["content"]


@router.post("", response_model=ChatResponse)
def chat(
    req: ChatRequest,
    db: Session = Depends(database.get_db),
    current_user: models.users = Depends(auth_utils.get_current_user),
):
    ids = set(req.material_ids)
    materials = (
        db.query(models.Material)
        .filter(models.Material.id.in_(ids), models.Material.owner_id == current_user.id)
        .all()
    )
    if len(materials) != len(ids):
        raise HTTPException(status_code=404, detail="One or more materials were not found.")

    msgs = [{"role": t.role, "content": t.content} for t in req.history[-20:]]
    while msgs and msgs[0]["role"] != "user":
        msgs.pop(0)
    msgs.append({"role": "user", "content": req.message})

    reply = ask_ollama(materials, msgs) if PROVIDER == "ollama" else ask_anthropic(materials, msgs)
    return {"reply": reply}
