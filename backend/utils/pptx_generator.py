import math
import re
import os
import sys

# Ensure project root is available for imports
PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

from utils.labels import get_labels

# Palette (matches the Draft Bench app)
PRIMARY, DEEP, SOFT, BORDER = "6D28D9", "2E1065", "F3EEFF", "E2D8F8"
TEXT, MUTED, WHITE, LAVENDER = "1E1033", "6B5E85", "FFFFFF", "DDD0FA"

SLIDE_W, SLIDE_H = 13.333, 7.5


def _rgb(hex_str: str) -> RGBColor:
    return RGBColor.from_string(hex_str)


def _fit_size(paragraphs, w_in, h_in, max_pt, min_pt, deva=False, gap=0.35):
    """Largest font size (pt) at which the paragraphs should fit the box. python-pptx can't measure text."""
    char_w = 0.62 if deva else 0.5
    usable_w = max((w_in - 0.3) * 72, 20)
    size = max_pt
    while size > min_pt:
        per_line = max(int(usable_w / (size * char_w)), 4)
        lines = sum(max(1, math.ceil(len(p) / per_line)) for p in paragraphs if p) or 1
        height = lines * size * 1.2 + max(len(paragraphs) - 1, 0) * size * gap
        if height <= (h_in - 0.2) * 72:
            break
        size -= 1
    return size


def _est_height(paragraphs, w_in, size, deva=False, gap=0.35):
    """Estimated height (inches) the paragraphs need at a given size."""
    char_w = 0.62 if deva else 0.5
    per_line = max(int(max((w_in - 0.3) * 72, 20) / (size * char_w)), 4)
    lines = sum(max(1, math.ceil(len(p) / per_line)) for p in paragraphs if p) or 1
    return (lines * size * 1.2 + max(len(paragraphs) - 1, 0) * size * gap) / 72 + 0.25


def _set_font(run, size, color, bold=False, italic=False, deva=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = _rgb(color)
    rPr = run._r.get_or_add_rPr()
    if deva:  # complex-script font so Devanagari renders with a proper font
        cs = etree.SubElement(rPr, qn("a:cs"))
        cs.set("typeface", "Nirmala UI")


def _bullet(paragraph, color):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", "342900")
    pPr.set("indent", "-342900")
    bu_clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(bu_clr, qn("a:srgbClr")).set("val", color)
    etree.SubElement(pPr, qn("a:buChar")).set("char", "•")


def _rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0, rounded=True, radius=0.08, alpha=None, shape=None):
    kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE)
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = _rgb(fill)
        if alpha is not None:
            clr = shp.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            etree.SubElement(clr, qn("a:alpha")).set("val", str(int(alpha * 1000)))
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = _rgb(line)
        shp.line.width = Pt(line_w)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    style = shp._element.find(qn("p:style"))  # drop the theme shadow in every viewer
    if style is not None:
        style.find(qn("a:effectRef")).set("idx", "0")
    return shp


def _text(slide, x, y, w, h, paragraphs, size, color=TEXT, bold=False, italic=False, align=PP_ALIGN.LEFT,
          anchor=MSO_ANCHOR.TOP, deva=False, bullets=False, bullet_color=PRIMARY, space_after=0.0):
    """Add a text box. paragraphs is a list of strings or of lists of (text, {bold, color, size}) runs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space_after:
            p.space_after = Pt(space_after)
        runs = para if isinstance(para, list) else [(para, {})]
        for txt, opt in runs:
            r = p.add_run()
            r.text = txt
            _set_font(r, opt.get("size", size), opt.get("color", color), opt.get("bold", bold), italic, deva)
        if bullets:
            _bullet(p, bullet_color)
    return tb


def _bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(color)


def _picture_in_frame(slide, img_path, x, y, w, h):
    """Place a picture centred inside a white rounded frame, keeping its proportions."""
    from PIL import Image
    _rect(slide, x, y, w, h, fill=WHITE, line=BORDER, line_w=1.5, radius=0.04)
    try:
        with Image.open(img_path) as im:
            iw, ih = im.size
        pad = 0.15
        bw, bh = w - 2 * pad, h - 2 * pad
        ratio = min(bw / iw, bh / ih)
        pw, ph = iw * ratio, ih * ratio
        slide.shapes.add_picture(img_path, Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))
        return True
    except Exception as e:
        print(f"Warning: PPTX image embed failed: {e}")
        return False


_NULLS = {"", "null", "none", "n/a", "na", "-"}
_JUNK_OUTCOME = re.compile(r"^(standard\s*\d+[,\s]*)+$", re.I)


def _clean(v) -> str:
    """Text for display; the model sometimes writes 'null' or 'None' as a string."""
    s = str(v if v is not None else "").strip()
    return "" if s.lower() in _NULLS else s


def _shorten(text: str, max_words: int = 10) -> str:
    """Safety net so one long sentence never breaks a slide."""
    words = text.split()
    if len(words) <= max_words:
        return text
    head = " ".join(words[: max_words + 2])
    cut = max(head.rfind(","), head.rfind(";"))
    if cut > len(head) * 0.4:
        return head[:cut].rstrip()
    return " ".join(words[:max_words]).rstrip(",;:") + "…"


def _notes(slide, info, L):
    """Everything that is useful for the teacher but would crowd the slide goes into the speaker notes."""
    parts = []
    if _clean(info.get("explanation")):
        parts.append(_clean(info["explanation"]))
    if _clean(info.get("example")):
        parts.append(f"{L.get('example', 'Example')}: {_clean(info['example'])}")
    terms = [t for t in (info.get("key_terms") or []) if isinstance(t, dict) and _clean(t.get("term"))]
    if terms:
        parts.append(f"{L.get('key_terms', 'Key words')}: " + "; ".join(
            f"{_clean(t['term'])} = {_clean(t.get('meaning'))}" for t in terms))
    for key, label in (("teacher_script", L["teacher_script"]), ("communication_support", L["communication"]),
                       ("sensory_note", L["sensory_note"])):
        if _clean(info.get(key)):
            parts.append(f"{label}: {_clean(info[key])}")
    if parts:
        slide.notes_slide.notes_text_frame.text = "\n\n".join(parts)


def _progress(slide, idx, total):
    """A quiet progress row: it shows where we are and how many slides are left (predictability)."""
    d, gap = 0.17, 0.33
    for i in range(1, total + 1):
        _rect(slide, 0.9 + (i - 1) * gap, 7.0, d, d, fill=PRIMARY if i <= idx else BORDER, shape=MSO_SHAPE.OVAL)
    _text(slide, SLIDE_W - 1.9, 6.9, 1.3, 0.35, [f"{idx} / {total}"], 12, color=MUTED, align=PP_ALIGN.RIGHT)


def _title_slide(prs, title, subtitle, meta_line, deva):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, PRIMARY)
    _rect(s, 8.6, -2.2, 6.6, 6.6, fill=WHITE, alpha=9, shape=MSO_SHAPE.OVAL)
    _rect(s, -1.6, 4.6, 4.6, 4.6, fill=WHITE, alpha=7, shape=MSO_SHAPE.OVAL)
    _rect(s, 0.9, 0.8, 0.9, 0.9, fill=WHITE, radius=0.22)
    _text(s, 0.9, 0.8, 0.9, 0.9, ["DB"], 26, color=PRIMARY, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    size = _fit_size([title], 10.5, 2.6, 54, 30, deva, gap=0)
    _text(s, 0.9, 2.5, 10.5, 2.6, [title], size, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE, deva=deva)
    _text(s, 0.9, 5.2, 10.5, 0.6, [subtitle], 24, color=LAVENDER, deva=deva)
    if meta_line:
        _text(s, 0.9, 5.85, 10.5, 0.5, [meta_line], 16, color=LAVENDER, deva=deva)
    return s


def _schedule_slide(prs, schedule, heading, L, idx, total, deva):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, WHITE)
    _rect(s, 0.6, 0.45, 0.12, 0.8, fill=PRIMARY, rounded=False)
    _text(s, 0.9, 0.4, 11, 0.9, [heading], 36, color=DEEP, bold=True, anchor=MSO_ANCHOR.MIDDLE, deva=deva)
    steps = schedule[:4]
    n = max(len(steps), 1)
    gap = 0.3
    cw = (SLIDE_W - 1.2 - gap * (n - 1)) / n
    for i, st in enumerate(steps):
        x = 0.6 + i * (cw + gap)
        _rect(s, x, 1.9, cw, 4.2, fill=SOFT, line=BORDER, radius=0.06)
        _rect(s, x + 0.25, 2.15, 0.7, 0.7, fill=PRIMARY, shape=MSO_SHAPE.OVAL)
        _text(s, x + 0.25, 2.15, 0.7, 0.7, [str(i + 1)], 24, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
              anchor=MSO_ANCHOR.MIDDLE)
        _text(s, x + 0.15, 3.05, cw - 0.3, 0.7, [_clean(st.get("label"))], 28, color=DEEP, bold=True, deva=deva)
        desc = _clean(st.get("icon_desc"))
        _text(s, x + 0.15, 3.8, cw - 0.3, 1.6, [desc], _fit_size([desc], cw - 0.3, 1.6, 22, 14, deva), color=TEXT, deva=deva)
        if st.get("duration_min"):
            _text(s, x + 0.15, 5.4, cw - 0.3, 0.5, [f"{st['duration_min']} {L['minutes']}"], 16, color=PRIMARY, bold=True, deva=deva)
    _progress(s, idx, total)
    return s


def _content_slide(prs, info, idx, total, L, deva):
    """One idea per slide: title, key idea, up to 3 short bullets, and one picture (or one example)."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, WHITE)
    title = _clean(info.get("title")) or f"{L['slide']} {idx}"
    key_idea = _shorten(_clean(info.get("key_idea")), 16)
    bullets = [_shorten(_clean(b), 10) for b in (info.get("bullets") or []) if _clean(b)][:3]
    example = _shorten(_clean(info.get("example")), 22)
    terms = [t for t in (info.get("key_terms") or []) if isinstance(t, dict) and _clean(t.get("term"))]
    img = info.get("image_path") or info.get("icon_path")
    has_img = bool(img and os.path.exists(img))

    right = "image" if has_img else ("example" if example else ("term" if terms else None))
    LX = 0.9
    LW = 7.4 if right else 11.6
    RX, RW = 8.7, 3.95

    _rect(s, 0.6, 0.45, 0.12, 0.85, fill=PRIMARY, rounded=False)
    tsize = _fit_size([title], 11.6, 1.0, 40, 28, deva, gap=0)
    _text(s, LX, 0.38, 11.8, 1.0, [title], tsize, color=DEEP, bold=True, anchor=MSO_ANCHOR.MIDDLE, deva=deva)

    y = 1.75
    if key_idea:
        _rect(s, LX, y, LW, 1.45, fill=SOFT, line=BORDER, radius=0.08)
        ksz = _fit_size([key_idea], LW - 0.4, 1.45, 32, 22, deva, gap=0)
        _text(s, LX + 0.2, y, LW - 0.4, 1.45, [key_idea], ksz, color=DEEP, bold=True, anchor=MSO_ANCHOR.MIDDLE, deva=deva)
        y += 1.75
    if bullets:
        bmax = 6.7 - y
        bsz = _fit_size(bullets, LW, bmax, 30, 22, deva, gap=0.5)
        bh = min(bmax, _est_height(bullets, LW, bsz, deva, 0.5))
        _text(s, LX, y, LW, bh, bullets, bsz, color=TEXT, bullets=True, deva=deva, space_after=bsz * 0.5)

    if right == "image":
        fh = 4.65
        if _picture_in_frame(s, img, RX, 1.75, RW, fh) and info.get("image_path") and info.get("image_page"):
            _text(s, RX, 1.75 + fh + 0.03, RW, 0.3, [f"{L.get('from_material', 'From your material')} · {L['page']} {info['image_page']}"],
                  11, color=MUTED, deva=deva)
    elif right in ("example", "term"):
        _rect(s, RX, 1.75, RW, 4.65, fill=PRIMARY if right == "example" else SOFT, line=None if right == "example" else BORDER, radius=0.06)
        if right == "example":
            _text(s, RX + 0.25, 1.95, RW - 0.5, 0.45, [L.get("example", "Example")], 16, color=LAVENDER, bold=True, deva=deva)
            _text(s, RX + 0.25, 2.5, RW - 0.5, 3.7, [example], _fit_size([example], RW - 0.5, 3.7, 30, 20, deva, gap=0),
                  color=WHITE, anchor=MSO_ANCHOR.MIDDLE, deva=deva)
        else:
            t = terms[0]
            meaning = _shorten(_clean(t.get("meaning")), 14)
            _text(s, RX + 0.25, 1.95, RW - 0.5, 0.45, [L.get("key_terms", "Key words")], 16, color=PRIMARY, bold=True, deva=deva)
            _text(s, RX + 0.25, 2.6, RW - 0.5, 1.2, [_clean(t["term"])], _fit_size([_clean(t["term"])], RW - 0.5, 1.2, 40, 26, deva, gap=0),
                  color=DEEP, bold=True, deva=deva)
            _text(s, RX + 0.25, 3.9, RW - 0.5, 2.3, [meaning], _fit_size([meaning], RW - 0.5, 2.3, 26, 18, deva, gap=0),
                  color=TEXT, deva=deva)

    _progress(s, idx, total)
    _notes(s, info, L)
    return s


def _done_slide(prs, idx, total, L, deva):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, SOFT)
    _rect(s, SLIDE_W / 2 - 1.2, 1.7, 2.4, 2.4, fill=PRIMARY, shape=MSO_SHAPE.OVAL)
    _text(s, SLIDE_W / 2 - 4, 4.4, 8, 1.4, [L.get("all_done", "All done!")], 54, color=DEEP, bold=True,
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, deva=deva)
    _progress(s, idx, total)
    return s


def generate_styled_pptx(
    slides_data: list,
    output_path: str,
    grade: int,
    subject: str,
    image_paths: list = None,
    language: str = "english",
    lesson_title: str = None,
    visual_schedule: list = None,
    cdc_outcome: str = None,
):
    L = get_labels(language)
    deva = (language or "").lower() == "nepali"
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    # Backward compatible: an aligned list of pictures, used only where a slide has no image of its own
    for i, info in enumerate(slides_data):
        if not info.get("image_path") and image_paths and i < len(image_paths) and image_paths[i]:
            info["image_path"] = image_paths[i]

    schedule = [st for st in (visual_schedule or []) if isinstance(st, dict) and _clean(st.get("label"))]
    outcome = _clean(cdc_outcome)
    if _JUNK_OUTCOME.match(outcome):
        outcome = ""
    total = 1 + (1 if schedule else 0) + len(slides_data) + 1

    _title_slide(prs, _clean(lesson_title) or subject, f"{L['grade']} {grade} • {subject}", outcome, deva)
    idx = 2
    if schedule:
        _schedule_slide(prs, schedule, L["visual_schedule"], L, idx, total, deva)
        idx += 1
    for info in slides_data:
        _content_slide(prs, info, idx, total, L, deva)
        idx += 1
    _done_slide(prs, idx, total, L, deva)

    prs.save(output_path)
    return output_path
